#!/usr/bin/env python3
"""Build the Sampling Sigma manuscript from the official IEEE Access Word template."""
from __future__ import annotations

import copy
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "2026-10-29"
RESULTS = json.loads((OUT / "evaluation" / "actual-prototype-results.json").read_text())
R = RESULTS["actual_results"]
IEEE_BLUE = RGBColor(0x00, 0x70, 0xA8)


def set_columns(section, count: int, space_twips: int = 360) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols"); sect_pr.append(cols)
    cols.set(qn("w:num"), str(count)); cols.set(qn("w:space"), str(space_twips))


def clear_body_keep_section(doc: Document) -> None:
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def p(doc, text="", style="PARA", bold_lead=None):
    para = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        para.add_run(bold_lead).bold = True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    return para


def heading(doc, number, title):
    para = doc.add_paragraph(style="H1")
    para.paragraph_format.keep_with_next = True
    para.add_run(f"{number}. {title.upper()}")
    return para


def subheading(doc, letter, title):
    para = doc.add_paragraph(style="H2_First")
    para.paragraph_format.keep_with_next = True
    # Disable the template's document-wide automatic letter sequence so each
    # IEEE section can restart its own A, B, C subsection lettering.
    num_pr = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "0")
    num_pr.append(num_id)
    para._p.get_or_add_pPr().append(num_pr)
    para.add_run(f"{letter}.    {title.upper()}")
    return para


def bullet(doc, text):
    para = doc.add_paragraph(style="PARA")
    para.paragraph_format.left_indent = Inches(.16)
    para.paragraph_format.first_line_indent = Inches(-.12)
    para.add_run("• "); para.add_run(text)


def caption(doc, number, title):
    para = doc.add_paragraph(style="Table Caption")
    para.paragraph_format.keep_with_next = True
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(f"TABLE {number}\n{title.upper()}")
    run.bold = False


def compact_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Normal Table"
    table.autofit = False
    if widths:
        for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
            grid_col.w = Inches(width)
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW"); tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
        tbl_w.set(qn("w:type"), "dxa")
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]; cell.text = h
        if widths:
            cell.width = Inches(widths[i])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(int(widths[i] * 1440))); tc_w.set(qn("w:type"), "dxa")
        for run in cell.paragraphs[0].runs:
            run.bold = True; run.font.size = Pt(8)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if widths:
                cells[i].width = Inches(widths[i])
                tc_w = cells[i]._tc.get_or_add_tcPr().find(qn("w:tcW"))
                tc_w.set(qn("w:w"), str(int(widths[i] * 1440))); tc_w.set(qn("w:type"), "dxa")
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs: run.font.size = Pt(8)
        cells[0]._tc.getparent().get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    return table


def reference(doc, number, text):
    para = doc.add_paragraph(style="Ref")
    para.add_run(f"[{number}] {text}")


def equation(doc, text, number):
    """Insert a native, editable Word math paragraph with an equation number."""
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.9)
    table.columns[1].width = Inches(.35)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    table._tbl.tblPr.append(borders)
    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = text
    math_run.append(math_text); math.append(math_run); math_para.append(math)
    left._p.append(math_para)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.add_run(f"({number})").font.size = Pt(7)
    return table


def algorithm(doc, title, lines):
    para = doc.add_paragraph(style="Table Caption")
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run(f"ALGORITHM 1\n{title.upper()}")
    table = doc.add_table(rows=1, cols=1)
    table.style = "Normal Table"
    cell = table.cell(0, 0)
    cell.text = ""
    for index, line in enumerate(lines):
        para = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.left_indent = Inches(.03)
        run = para.add_run(line)
        run.font.name = "Courier New"; run.font.size = Pt(7.5)
    return table


def build(template: Path, output: Path) -> None:
    doc = Document(template)
    doc.core_properties.author = "Sri Sai Lakshmi Ganesh Murukutla"
    doc.core_properties.last_modified_by = "Sri Sai Lakshmi Ganesh Murukutla"
    doc.core_properties.title = "Scope-Bound Cybersecurity Evidence Passports With Independent Review and Living Assurance Lineage"
    doc.core_properties.subject = "Sampling Sigma cybersecurity evidence-assurance research"
    doc.core_properties.keywords = "cybersecurity assurance, evidence provenance, independent review, verifiable credentials, smart contracts"
    original_sect_pr = doc.sections[0]._sectPr
    recurring_refs = [copy.deepcopy(node) for node in original_sect_pr if node.tag in {qn("w:headerReference"), qn("w:footerReference")}]
    clear_body_keep_section(doc)
    first = doc.sections[0]
    first.page_width = Inches(8); first.page_height = Inches(10.88)
    first.left_margin = Inches(.51); first.right_margin = Inches(.51)
    first.top_margin = Inches(.89); first.bottom_margin = Inches(.72)
    set_columns(first, 1)

    p(doc, "Date of publication xxxx 00, 0000, date of current version xxxx 00, 0000.", "DOP")
    p(doc, "Digital Object Identifier: To be assigned after acceptance", "DOI")
    title = doc.add_paragraph(style="Paper Title")
    title.add_run("Scope-Bound Cybersecurity Evidence Passports With Independent Review and Living Assurance Lineage")
    p(doc, "Sri Sai Lakshmi Ganesh Murukutla, Murukutla Group LLC", "AU")
    p(doc, "Murukutla Group LLC, Austin, Texas, USA", "Aff")
    p(doc, "Corresponding author: Sri Sai Lakshmi Ganesh Murukutla (email: murukutlassl@gmail.com). ORCID: https://orcid.org/0009-0001-5057-7399.", "CA")
    p(doc, "Funding information: This research received no external funding.", "PI")

    abstract_text = (
        "Organizations routinely exchange cybersecurity reports, attestations, scan results, and provenance records. Once separated from the processes that produced them, however, these records may retain a valid digest while losing the context needed to judge scope, exclusions, reviewer independence, currency, and supersession. "
        "This study develops Sampling Sigma, a design science prototype for scope-bound cybersecurity evidence passports. Confidential evidence remains off-chain; a portable passport binds the artifact digest to its subject, declared scope, limitations, validity period, issuer, independent review, conflict disclosure, lifecycle status, and lineage. "
        "The prototype comprises a Solidity contract, a local verifier, and a reproducible evaluation harness. All 15 contract tests passed. In controlled experiments, the verifier detected every one of 5,000 seeded single-bit mutations, and the reference lifecycle model rejected all 21 forbidden transitions among 25 ordered state pairs. "
        "These results demonstrate internal consistency under the reported conditions. They do not establish semantic truth, security certification, interoperability, usability, or regulatory compliance. The study therefore separates integrity, provenance, validation, and economics and specifies a preregisterable comparison of raw digest exchange, signed centralized records, credential-only passports, and an optional shared registry."
    )
    para = doc.add_paragraph(style="Abstract")
    label = para.add_run("ABSTRACT "); label.bold = True; label.font.color.rgb = IEEE_BLUE
    para.add_run(abstract_text)
    para = doc.add_paragraph(style="IT")
    label = para.add_run("INDEX TERMS "); label.bold = True; label.font.color.rgb = IEEE_BLUE
    para.add_run("Blockchain, cybersecurity assurance, design science, evidence provenance, independent review, living assurance, privacy, smart contracts, verifiable credentials.")

    second = doc.add_section(WD_SECTION.CONTINUOUS)
    second.page_width = Inches(8); second.page_height = Inches(10.88)
    second.left_margin = Inches(.51); second.right_margin = Inches(.51)
    second.top_margin = Inches(.90); second.bottom_margin = Inches(.72)
    second.header.is_linked_to_previous = True; second.footer.is_linked_to_previous = True
    set_columns(second, 2)

    heading(doc, "I", "Introduction")
    p(doc, "Cybersecurity assurance depends on more than preserving a file's bytes. A relying party must also know what was assessed, what remained outside the assessment, who approved the evidence, whether the reviewer was independent, and whether a later record superseded the one presented. Secure timestamping [1], supply-chain attestation [2], verifiable credentials [3], and cybersecurity-control assessment [4] address portions of this problem. Used separately, however, they do not provide a compact decision record that carries integrity, scope, limitations, review, validity, and lineage together.")
    p(doc, "Sampling Sigma examines whether a portable evidence passport can preserve that decision context without transferring the underlying confidential artifact. The design deliberately treats a shared ledger as optional. Its inclusion must be supported by evidence that cross-organizational coordination improves beyond what a signed centralized record can provide.")
    p(doc, "The work makes four contributions. First, it defines a minimum passport schema that records both assurance claims and their boundaries. Second, it separates the issuer, reviewer organization, relying party, and economic roles. Third, it models rejection, expiry, revocation, and supersession as explicit lifecycle events. Finally, it couples a reproducible prototype evaluation with a preregisterable protocol for subsequent human-subject testing.")

    heading(doc, "II", "Research Questions and Related Work")
    subheading(doc, "A", "Research Questions")
    bullet(doc, "RQ1: Which fields prevent a fingerprint from being mistaken for proof of truth or full-scope assurance?")
    bullet(doc, "RQ2: Can organizational separation and conflict commitments make silent self-approval detectable?")
    bullet(doc, "RQ3: Can explicit validity and lineage reduce acceptance of stale or superseded evidence?")
    bullet(doc, "RQ4: Can local verification preserve confidentiality while remaining usable?")
    bullet(doc, "RQ5: Under which multi-organization conditions does a shared registry add value over signed centralized records?")
    subheading(doc, "B", "Research Gap")
    p(doc, "Existing standards supply many of the required building blocks. Haber and Stornetta established tamper-evident document timestamping [1]. in-toto provides authenticated software-supply-chain provenance [2], and SLSA defines progressive artifact-integrity practices [5]. The W3C Verifiable Credentials Data Model contributes portable claims, issuer proofs, and status semantics [3]. NIST SP 800-53A frames evidence-driven control assessment [4], while C2PA offers a related model of provenance and trust [6]. What remains underdeveloped is a compact record that assembles these capabilities into enough decision context for a relying party to judge a cybersecurity assurance claim without receiving the confidential source artifact.")

    heading(doc, "III", "Design Science Method")
    p(doc, "The study follows the design science cycle described by Hevner et al. [7], progressing from problem identification and requirements derivation through construction, demonstration, evaluation, and communication. Requirements were synthesized from guidance on assessment, secure development, provenance, credentials, and threat modeling [3]-[6], [8]. Evaluation proceeds in three stages: deterministic contract tests, controlled mutation and reference state-machine experiments, and a planned comparative human study. Measured results are currently available for the first two stages only. Table I summarizes the resulting requirements and protected failure modes.")
    caption(doc, "I", "Design requirements and protected failure modes")
    compact_table(doc, ["ID", "Requirement", "Failure"], [
        ["R1", "Bind artifact, subject, scope, limitations, and period", "Context collapse"],
        ["R2", "Separate issuer and reviewer organization", "Self-approval"],
        ["R3", "Commit conflicts and review evidence", "Hidden dependency"],
        ["R4", "Expose validity, rejection, revocation, supersession", "Stale assurance"],
        ["R5", "Keep evidence off-chain; verify locally", "Disclosure"],
        ["R6", "Separate validity from token economics", "Economic capture"],
    ], [.35, 1.75, 1.15])

    heading(doc, "IV", "Architecture and Data Model")
    subheading(doc, "A", "Assurance Planes")
    p(doc, "The architecture distributes responsibility across five assurance planes and an optional economic plane. Source artifacts remain in an encrypted evidence vault. The passport service canonicalizes, hashes, validates, and signs the portable record, while an independence graph represents organizations and disclosed relationships. An assurance registry stores minimum commitments and lifecycle events. At the point of use, the local verifier recalculates the digest, validates the signature and status, checks validity and supersession, and applies the relying party's policy. The optional economic layer is structurally excluded from decisions about evidentiary truth. Table II groups the fields carried by the evidence passport.")
    caption(doc, "II", "Evidence passport field groups")
    compact_table(doc, ["Group", "Required fields"], [
        ["Identity", "Schema version, passport ID, subject commitment"],
        ["Integrity", "Algorithm, artifact digest, canonicalization profile"],
        ["Context", "Domain, assessment type, scope and limitation commitments"],
        ["Review", "Issuer, reviewer organization, conflict and evidence commitments"],
        ["Lifecycle", "Issued time, valid-until, status, revocation, supersession"],
        ["Proof", "Signature suite, key identifier, policy and status endpoints"],
    ], [.75, 2.55])
    p(doc, "A successful digest check establishes byte equality with the committed artifact. Correctness, completeness, authorization, reviewer competence, and continuing security require separate evidence. The verifier states this boundary during the decision process so that a valid digest cannot be presented as broader proof.")
    subheading(doc, "B", "Formal Commitments and Decision Predicate")
    p(doc, "Let x be a JSON evidence passport excluding its proof. The portable commitment uses the JSON Canonicalization Scheme [12] followed by SHA-256 as specified by FIPS 180-4 [11].")
    equation(doc, "C(x) = SHA-256(JCS(x))", 1)
    p(doc, "For passport P, artifact A, and decision time t, acceptance is conjunctive rather than score-based. Each indicator equals one only when its named check passes.")
    equation(doc, "V = H ∧ S ∧ C ∧ R ∧ T ∧ Q", 2)
    p(doc, "H recomputes the artifact commitment; S validates the issuer proof; C requires the relying party to acknowledge scope and limitations; R checks reviewer identity, organizational separation, and conflict disclosure; T rejects expired, revoked, rejected, or superseded passports; and Q applies local policy. A zero in any term produces rejection. This model prevents confidence scores from overriding a failed mandatory control. Table III summarizes the threats used to challenge these checks.")

    heading(doc, "V", "Threat Model")
    caption(doc, "III", "Threats, mitigations, and residual risks")
    compact_table(doc, ["Threat", "Mitigation", "Residual risk"], [
        ["Substitution", "Local digest and signature check", "Compromised endpoint"],
        ["Scope inflation", "Mandatory scope and limitations", "Deceptive wording"],
        ["Self-review", "Role and organization separation", "Collusion"],
        ["Stale evidence", "Expiry, revocation, supersession", "Publication delay"],
        ["Data leakage", "Off-chain evidence and minimization", "Metadata correlation"],
        ["Admin compromise", "Role separation, multisignature target, audit", "Key compromise"],
    ], [.8, 1.45, 1.05])

    heading(doc, "VI", "Prototype Implementation")
    subheading(doc, "A", "Smart Contract")
    p(doc, "The prototype uses Solidity 0.8.28 and OpenZeppelin components [9]. SamplingSigmaAssuranceCore records subject, scope, limitations, validity, reviewer organization, conflict commitment, rejection, revocation, and supersession. A browser-local verifier hashes files and produces a JSON passport without uploading the source. The separate token candidate is not required for evidence verification, and no public sale, exchange listing, or production deployment is authorized by this research artifact.")
    subheading(doc, "B", "Verification Algorithm")
    algorithm(doc, "Local passport verification", [
        "INPUT: passport P, local artifact A, time t, policy Q",
        "1  require schema(P) and canonicalizationProfile(P)",
        "2  require SHA256(A) = P.artifactHash",
        "3  require verifySignature(P.proof, issuerKey(P))",
        "4  require scopeAccepted(P.scope, P.limitations)",
        "5  require independentReviewer(P) and conflictChecked(P)",
        "6  require P.status = VALIDATED and t <= P.validUntil",
        "7  require notRevoked(P) and notSuperseded(P)",
        "8  require Q(P) = true; return ACCEPT",
        "9  on any failed requirement: return REJECT with reason",
    ])
    p(doc, "The verifier applies fail-closed semantics. Malformed records, unknown algorithms, missing commitments, failed proofs, unavailable status information, and unmet policy conditions each produce a rejection with a stated reason. The implementation uses explicit custom errors and emits lifecycle events to preserve an auditable decision trail.")
    subheading(doc, "C", "Reproducibility Package")
    p(doc, "The reproducibility package records the source revision, dependency lock file, test output, evaluation script, raw JSON results, machine environment, and document checksum. The public research artifact is available at https://github.com/SriMurukutla/sampling-sigma-evidence-passports. The evaluated submission snapshot is anchored by commit 3df129504776c6313f0b4a03ec0f242caa88a94d; a separate archival DOI may be added after depositing the tagged release in a long-term repository.")

    heading(doc, "VII", "Measured Prototype Results")
    subheading(doc, "A", "Experimental Setup")
    p(doc, "Study SS-DSR-PROTOTYPE-2026-10-29 executed three deterministic experiments and the repository contract suite. The mutation experiment generated 5,000 seeded byte strings, flipped one indexed bit in each string, and compared SHA-256 digests. The lifecycle reference model enumerated all 25 ordered pairs over Submitted, Validated, Rejected, Revoked, and Superseded. The performance test hashed a deterministic one-mebibyte payload 256 times using a monotonic nanosecond timer. Contract tests ran through the repository's npm test command. Tables IV and V report the observed measurements and contract-suite composition.")
    subheading(doc, "B", "Observed Results")
    caption(doc, "IV", "Observed prototype measurements")
    t=R["tamper_detection"]; h=R["hashing_performance"]; l=R["lifecycle_reference_model"]; c=R["smart_contract_tests"]
    compact_table(doc, ["Experiment", "Observed result", "Boundary"], [
        ["Contract suite", f"{c['passing_tests']} passing; exit {c['exit_code']}", "Not an audit"],
        ["Mutation corpus", f"{t['single_bit_mutations_detected']}/{t['cases']} detected", "Byte changes only"],
        ["Lifecycle model", f"{l['state_pairs_exhaustively_checked']} pairs; {l['forbidden_transitions_rejected']} forbidden rejected", "Reference model"],
        ["SHA-256 benchmark", f"Median {h['median_latency_ms']} ms; p95 {h['p95_latency_ms']} ms", "Single machine/run"],
    ], [1.0, 1.25, 1.05])
    p(doc, f"The controlled corpus detected 100% of 5,000 injected single-bit mutations. Under a simplifying independent-Bernoulli model, the exact two-sided 95% lower confidence bound for detection in this generator is 99.9263%.")
    equation(doc, "L95 = 0.025^(1/5000) = 0.999263", 3)
    p(doc, f"The bound describes only this seeded mutation generator; it does not estimate SHA-256 collision resistance, validate semantic truth, or establish an operational false-acceptance probability. The benchmark used Python {RESULTS['environment']['python']} on {RESULTS['environment']['platform']} and measured {h['throughput_mib_per_second']} MiB/s for 256 one-mebibyte payloads. The latency distribution is reported by median and p95 because no normality assumption was made.")
    caption(doc, "V", "Contract suite composition")
    compact_table(doc, ["Component", "Tests", "Primary behaviors"], [
        ["Assurance core", "4", "Commitments, independent review, supersession"],
        ["Research core", "6", "Evidence validation, authorization, immutability"],
        ["Token candidate", "5", "Fixed allocation, role separation, pause controls"],
    ], [1.0, .45, 1.85])
    p(doc, "All 15 repository tests passed in the recorded run. Four tests exercised the assurance core, six addressed research-record governance, and five covered the separate token candidate. The token results are included to document repository completeness; they do not answer the assurance research questions. The prototype has not yet undergone an independent audit, user study, interoperability assessment, production-availability study, or legal evaluation.")

    heading(doc, "VIII", "Comparative Evaluation Protocol")
    p(doc, "A controlled study will compare four conditions: raw digest exchange; a signed centralized JSON record; a W3C Verifiable Credential passport without a ledger; and Sampling Sigma with an optional shared registry. Participants will issue and verify identical evidence packages containing altered, expired, superseded, conflicted, and out-of-scope cases. An exploratory pilot of 24 to 30 participants will test the instruments and estimate variance. A confirmatory sample will then be determined by preregistered power analysis using a declared smallest effect of interest. Table VI defines the planned outcomes and analyses.")
    caption(doc, "VI", "Planned primary outcomes and analysis")
    compact_table(doc, ["Outcome", "Definition", "Analysis"], [
        ["Accuracy", "Correct accept/reject decision", "Mixed logistic model"],
        ["Time", "Seconds to defensible decision", "Survival/log-normal model"],
        ["Critical error", "Invalid accepted or valid rejected", "Risk difference and CI"],
        ["Completeness", "Required context interpreted", "Rubric and agreement"],
        ["Privacy", "Evidence bytes transmitted", "Network capture"],
        ["Usability", "SUS and task workload", "CI and effect size"],
    ], [.85, 1.45, 1.0])
    p(doc, "The confirmatory protocol will fix primary outcomes, exclusions, and analysis before data collection. Holm correction will control multiplicity. Effect sizes and 95% confidence intervals will accompany p-values. Missing data, deviations, adverse events, and exclusions will be reported. Robust or nonparametric sensitivity analyses will address distributional violations.")

    heading(doc, "IX", "Prospective Results and Decision Rules")
    p(doc, "The following decision rules concern prospective outcomes and should not be read as observed findings. A shared registry will be retained only if the comparative study identifies a measurable advantage over simpler designs.")
    bullet(doc, "If the lower 95% confidence limit for decision-accuracy improvement is above zero without increased critical errors, proceed to a larger pilot.")
    bullet(doc, "If the credential-only passport is within the preregistered equivalence margin, prefer that simpler architecture.")
    bullet(doc, "If users over-trust successful digest checks and the critical misconception rate exceeds 5%, redesign the verifier language and workflow.")
    bullet(doc, "If reviewer independence cannot be verified, do not claim independent-review assurance.")
    bullet(doc, "If metadata-correlation risk exceeds the approved privacy threshold, remove or reduce public metadata.")

    heading(doc, "X", "Validity, Ethics, and Reproducibility")
    subheading(doc, "A", "Validity")
    p(doc, "Construct validity is protected by distinguishing digest equality from truth. Internal validity uses identical task packages, randomized order, logged deviations, and blinded scoring where practical. External validity is limited: results from a controlled pilot will not be generalized to all sectors or jurisdictions.")
    subheading(doc, "B", "Ethics and Privacy")
    p(doc, "No real customer secrets or personal data will enter a public registry. Human-participant research requires the applicable institutional determination, informed consent where required, minimization, retention limits, and a withdrawal process. Security findings will follow coordinated disclosure.")
    subheading(doc, "C", "Reproducibility")
    p(doc, "The release will include schemas, fixtures, source code, dependency versions, raw measurements, test output, analysis scripts, and replication instructions where legally permitted. IEEE's Code Ocean and DataPort integrations provide possible archival channels [10].")

    heading(doc, "XI", "Governance and Economic Separation")
    p(doc, "The assurance protocol can operate without a tradable token. Any token distribution, sale, custody, listing, compensation, staking, or investment-oriented communication requires separate legal, tax, sanctions, consumer-protection, and securities analysis. Evidence validity must never depend on token ownership, price, or voting power. The article evaluates assurance architecture and does not make an investment claim.")

    heading(doc, "XII", "Limitations")
    bullet(doc, "The measurements are local prototype results and have not been independently replicated.")
    bullet(doc, "The lifecycle experiment checks a reference model and is not a formal contract-equivalence proof.")
    bullet(doc, "Identity, reviewer competence, organizational independence, and conflicts require external validation.")
    bullet(doc, "No field evidence yet establishes decision improvement, usability, interoperability, or economic value.")
    bullet(doc, "Public registries can introduce permanence, cost, governance, and metadata-correlation risks.")

    heading(doc, "XIII", "Conclusion")
    p(doc, "Sampling Sigma provides a structured way to carry artifact integrity, decision context, independent review, and lifecycle state across organizational boundaries while keeping the source evidence under local control. The prototype behaved consistently in the reported contract, mutation, and state-transition tests. That result is encouraging but narrow. Claims about practical usefulness must await the planned comparison with simpler signed records and credential-only passports. Continued use of the shared-registry design should depend on demonstrated gains in decision quality, privacy, continuity, or cross-organizational coordination.")

    heading(doc, "ACKNOWLEDGMENT", "")
    p(doc, "OpenAI ChatGPT [13] and Codex [14] were used for editorial restructuring, software-development assistance, preparation of the reproducible evaluation harness, and document formatting in the Abstract, Sections I-XIII, and Appendix A. The author reviewed the manuscript, source code, citations, experiments, and submission statements and independently executed the reported tests. Generated text was not treated as empirical evidence.")

    heading(doc, "DATA AVAILABILITY", "")
    p(doc, "The prototype source, dependency lock file, complete contract tests, evaluation script, raw JSON measurements, test transcript, checksum manifest, and submission-candidate manuscript are publicly available at https://github.com/SriMurukutla/sampling-sigma-evidence-passports. The paper's evaluated snapshot is commit 3df129504776c6313f0b4a03ec0f242caa88a94d. The repository is public under the MIT License for code; no independent audit, production approval, or regulatory determination is implied. Table VII lists the minimum reproducibility artifacts.")

    heading(doc, "APPENDIX A", "Reproducibility Artifacts")
    caption(doc, "VII", "Reproducibility artifacts")
    compact_table(doc, ["Artifact", "Purpose"], [
        ["contracts/SamplingSigmaAssuranceCore.sol", "Assurance record and lifecycle implementation"],
        ["test/SamplingSigmaAssuranceCore.ts", "Assurance-core behavioral tests"],
        ["tools/research_evaluation.py", "Mutation, lifecycle, benchmark, and test harness"],
        ["evaluation/actual-prototype-results.json", "Machine-readable observed measurements"],
        ["evaluation/hardhat-test-output.txt", "Recorded contract-test output"],
    ], [1.65, 1.65])
    p(doc, "Minimum replication sequence: install the locked dependencies; run npm test; run python tools/research_evaluation.py; compare the resulting JSON study identifier, environment, test exit code, case counts, and timing statistics with the reported artifact. Timing values are expected to vary by machine; logical counts and rejection outcomes should not.")

    heading(doc, "REFERENCES", "")
    references = [
        "S. Haber and W. S. Stornetta, ‘How to time-stamp a digital document,’ J. Cryptology, vol. 3, pp. 99-111, 1991, doi: 10.1007/BF00196791.",
        "L. Torres-Arias et al., ‘in-toto: Providing farm-to-table guarantees for bits and bytes,’ in Proc. 28th USENIX Security Symp., 2019, pp. 1393-1410.",
        "W3C, Verifiable Credentials Data Model v2.0, W3C Recommendation, May 2025. [Online]. Available: https://www.w3.org/TR/vc-data-model-2.0/",
        "NIST, Assessing Security and Privacy Controls in Information Systems and Organizations, NIST SP 800-53A Rev. 5, Jan. 2022, doi: 10.6028/NIST.SP.800-53Ar5.",
        "SLSA, ‘Supply-chain Levels for Software Artifacts Specification v1.2.’ [Online]. Available: https://slsa.dev/spec/v1.2/",
        "C2PA, ‘Content Credentials Technical Specification, version 2.5.’ [Online]. Available: https://spec.c2pa.org/specifications/specifications/2.5/index.html",
        "A. R. Hevner, S. T. March, J. Park, and S. Ram, ‘Design science in information systems research,’ MIS Quart., vol. 28, no. 1, pp. 75-105, 2004, doi: 10.2307/25148625.",
        "NIST, Secure Software Development Framework, NIST SP 800-218, Feb. 2022, doi: 10.6028/NIST.SP.800-218.",
        "OpenZeppelin, ‘OpenZeppelin Contracts documentation.’ [Online]. Available: https://docs.openzeppelin.com/contracts/",
        "IEEE Access, ‘Preparing your article.’ [Online]. Available: https://ieeeaccess.ieee.org/authors/preparing-your-article/",
        "National Institute of Standards and Technology, Secure Hash Standard, FIPS PUB 180-4, Aug. 2015, doi: 10.6028/NIST.FIPS.180-4.",
        "A. Rundgren, B. Jordan, and S. Erdtman, ‘JSON Canonicalization Scheme (JCS),’ RFC 8785, Jun. 2020, doi: 10.17487/RFC8785.",
        "OpenAI, ‘ChatGPT.’ [Online]. Available: https://chatgpt.com/ (accessed Oct. 29, 2026).",
        "OpenAI, ‘Codex.’ [Online]. Available: https://openai.com/codex/ (accessed Oct. 29, 2026).",
    ]
    for i, item in enumerate(references, 1): reference(doc, i, item)

    heading(doc, "AUTHOR BIOGRAPHY", "")
    bio = doc.add_paragraph(style="AU_Bios")
    bio.add_run("SRI SAI LAKSHMI GANESH MURUKUTLA ").bold = True
    bio.add_run("is the founder of Murukutla Group LLC. His current work examines cybersecurity evidence provenance, independent assurance, privacy-preserving verification, secure software development, and accountable distributed-system architecture. ORCID: 0009-0001-5057-7399.")

    for section in doc.sections:
        sect_pr = section._sectPr
        for node in list(sect_pr):
            if node.tag in {qn("w:headerReference"), qn("w:footerReference")}:
                sect_pr.remove(node)
        for node in recurring_refs:
            sect_pr.insert(0, copy.deepcopy(node))
    for section in doc.sections:
        for para in section.footer.paragraphs:
            if "VOLUME XX, 2017" in para.text:
                for run in para.runs:
                    if "VOLUME XX, 2017" in run.text:
                        run.text = run.text.replace("VOLUME XX, 2017", "VOLUME XX, 2026")
            # The downloaded IEEE template stores the sample page number as
            # literal text. Replace it with a real PAGE field so the footer
            # increments correctly after the manuscript is repaginated.
            for run in para.runs:
                if run.text.strip() == "1":
                    for text_node in run._r.findall(qn("w:t")):
                        run._r.remove(text_node)
                    begin = OxmlElement("w:fldChar")
                    begin.set(qn("w:fldCharType"), "begin")
                    instruction = OxmlElement("w:instrText")
                    instruction.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    instruction.text = " PAGE "
                    end = OxmlElement("w:fldChar")
                    end.set(qn("w:fldCharType"), "end")
                    for node in (begin, instruction, end):
                        run._r.append(node)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_ieee_access_manuscript.py TEMPLATE.docx OUTPUT.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]))
